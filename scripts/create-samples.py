"""Build synthetic, non-personal test assets. Requires Pillow and openpyxl.
Usage: python scripts/create-samples.py <HWPX base template directory>
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, ZIP_STORED
import sys, re, html
from openpyxl import Workbook
from PIL import Image, ImageDraw

root=Path(__file__).resolve().parents[1]
base=Path(sys.argv[1])
out=root/'assets'/'samples'
out.mkdir(parents=True,exist_ok=True)
original=(base/'Contents/section0.xml').read_text(encoding='utf-8')
opening=original[:original.index('<hp:p')]
secpr=re.search(r'<hp:secPr\b.*?</hp:secPr>',original,re.S).group()
colpr=re.search(r'<hp:ctrl>.*?</hp:ctrl>',original,re.S).group()
def hwpx(name,lines):
    paragraphs=[]
    for i,line in enumerate(lines):
        structural=secpr+colpr if i==0 else ''
        paragraphs.append(f'<hp:p id="{100+i}" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0"><hp:run charPrIDRef="0">{structural}<hp:t>{html.escape(line)}</hp:t></hp:run></hp:p>')
    section=opening+''.join(paragraphs)+'</hs:sec>'
    with ZipFile(out/name,'w',ZIP_DEFLATED) as z:
        z.writestr('mimetype','application/hwp+zip',compress_type=ZIP_STORED)
        for p in base.rglob('*'):
            if not p.is_file() or p.name=='mimetype': continue
            key=p.relative_to(base).as_posix()
            if key=='Contents/section0.xml': z.writestr(key,section)
            elif key=='Preview/PrvText.txt': z.writestr(key,'\n'.join(lines))
            else: z.write(p,key)
hwpx('01_가을체험학습_계획서.hwpx',[
 '2026 가을 현장체험학습 운영 계획 (가상 예제)',
 '행사일: 2026. 10. 23. (금)',
 '장소: 푸른숲 생태공원',
 '참가 인원: 118명',
 '1인당 단가: 25,000원',
 '총예산: 2,950,000원',
 '제출 마감일: 2026. 10. 16. (금)',
 '이 문서는 앱 체험을 위한 가상 자료이며 실제 학생 정보는 없습니다.' ])
hwpx('02_가을체험학습_가정통신문.hwpx',[
 '2026 가을 현장체험학습 안내 (가상 예제)',
 '행사일: 2026. 10. 23. (목)',
 '장소: 푸른숲 생태공원',
 '참가 인원: 120명',
 '1인당 단가: 25,000원',
 '총예산: 3,000,000원',
 '제출 마감일: 2025. 10. 16.',
 '계획서와 다른 인원·예산, 잘못된 요일, 지난 연도 마감일을 찾아보세요.' ])
wb=Workbook();ws=wb.active;ws.title='체험학습 예산'
for row in [['2026 가을 체험학습 예산 (가상 예제)'],['장소','푸른숲 생태공원'],['참가 인원',120],['1인당 단가',25000],['총예산',3000000],[],['항목','수량','단가','금액'],['체험비',118,25000,3000000]]:ws.append(row)
from openpyxl.styles import Font, PatternFill, Alignment
for row in ws:
    for c in row:c.font=Font(name='맑은 고딕',size=11);c.alignment=Alignment(vertical='center');ws.row_dimensions[c.row].height=25
for row in [1,7]:
    for c in ws[row]:c.fill=PatternFill('solid',fgColor='236959');c.font=Font(name='맑은 고딕',color='FFFFFF',bold=True)
for col,width in [('A',42),('B',24),('C',20),('D',22)]:ws.column_dimensions[col].width=width
wb.save(out/'03_가을체험학습_예산표.xlsx')
# Original vector-like mark rendered with Pillow, no third-party icon asset.
im=Image.new('RGBA',(256,256),(0,0,0,0));d=ImageDraw.Draw(im)
d.rounded_rectangle((4,4,252,252),radius=55,fill='#236959')
d.rounded_rectangle((49,91,89,194),radius=15,fill='#a4c397')
d.rounded_rectangle((104,53,144,194),radius=15,fill='#ffffff')
d.rounded_rectangle((159,103,199,194),radius=15,fill='#d3dfa3')
im.save(root/'assets'/'icon.ico',sizes=[(16,16),(32,32),(48,48),(64,64),(128,128),(256,256)])
im.save(root/'assets'/'icon.png')
print('Created 3 synthetic documents and app icon.')
